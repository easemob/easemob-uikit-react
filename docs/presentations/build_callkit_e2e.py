"""Build the CallKit E2E sharing deck (editable PowerPoint/Keynote import)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("callkit-e2e-sharing.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = "101A31"
NAVY_2 = "172641"
TEAL = "22C7B4"
BLUE = "5B8DEF"
AMBER = "F4B759"
RED = "EC6471"
WHITE = "FFFFFF"
PAPER = "F5F7FB"
INK = "19243A"
MUTED = "62718A"
LINE = "DCE4EE"
FONT = "PingFang SC"


def rgb(hexcode):
    return RGBColor.from_string(hexcode)


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.02):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for idx, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(0)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
    return box


def line(slide, x1, y1, x2, y2, color=LINE, width=1.4):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                       Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def circle(slide, x, y, d, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    return shape


def base(title, kicker, number, dark=False, time=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = NAVY if dark else PAPER
    rect(slide, 0, 0, 13.333, 7.5, bg)
    rect(slide, 0, 0, 0.12, 7.5, TEAL)
    text(slide, kicker.upper(), 0.62, 0.35, 10.8, 0.28, 10,
         TEAL if dark else BLUE, True)
    text(slide, title, 0.62, 0.82, 12.1, 0.72, 29,
         WHITE if dark else INK, True)
    line(slide, 0.62, 6.99, 12.72, 6.99, "35425A" if dark else LINE)
    text(slide, "CALLKIT · E2E 测试实践", 0.65, 7.08, 5.5, 0.19, 9,
         "8C9AB3" if dark else MUTED)
    text(slide, f"{number:02d} / 09", 11.9, 7.08, 0.8, 0.19, 9,
         "8C9AB3" if dark else MUTED, align=PP_ALIGN.RIGHT)
    if time:
        text(slide, time, 11.14, 0.4, 1.55, 0.27, 10,
             "8C9AB3" if dark else MUTED, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, title, body, accent=TEAL, dark=False):
    rect(slide, x, y, w, h, NAVY_2 if dark else WHITE, True)
    rect(slide, x, y, 0.08, h, accent)
    text(slide, title, x+0.27, y+0.22, w-0.5, 0.45, 19,
         WHITE if dark else INK, True)
    text(slide, body, x+0.27, y+0.82, w-0.52, h-1.03, 15,
         "B9C5D8" if dark else MUTED, valign=MSO_ANCHOR.TOP)


# 01 — title
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0, 0, 0.16, 7.5, TEAL)
circle(s, 9.55, 1.0, 2.8, NAVY_2)
circle(s, 10.22, 1.67, 1.45, TEAL)
circle(s, 10.72, 2.17, 0.45, WHITE)
line(s, 10.95, 3.13, 10.95, 5.14, TEAL, 5)
text(s, "CALLKIT / E2E", 0.75, 1.0, 6.5, 0.36, 15, TEAL, True)
text(s, "把通话测到\n真的听见、真的看见", 0.75, 1.75, 9.1, 2.0, 36, WHITE, True)
text(s, "从信令状态，到多设备归属，再到真实 RTC 媒体内容", 0.79, 4.31, 9.3, 0.53,
     19, "B9C5D8")
rect(s, 0.78, 5.68, 4.33, 0.58, NAVY_2, True)
text(s, "团队分享  ·  10–15 分钟", 1.04, 5.8, 3.81, 0.28, 14, WHITE)
text(s, "01 / 09", 11.83, 7.05, 0.8, 0.22, 10, "8C9AB3", align=PP_ALIGN.RIGHT)

# 02 — challenge
s = base("一次通话成功，到底意味着什么？", "01 / 测试目标", 2, time="约 1 分钟")
text(s, "点击“接听”只是开始。我们要同时证明四件事。", 0.65, 1.68, 12.0, 0.48,
     19, MUTED)
labels = [
    ("01", "信令", "创建 / 接听 / 挂断的服务端状态正确", BLUE),
    ("02", "设备", "同账号多端只有正确设备进入通话", TEAL),
    ("03", "界面", "邀请、通话中、结束状态与快照一致", AMBER),
    ("04", "媒体", "远端真的收到持续变化的音视频", RED),
]
for i, (n, title, body, accent) in enumerate(labels):
    x = 0.65 + i*3.12
    card(s, x, 2.65, 2.82, 2.72, title, body, accent)
    text(s, n, x+2.02, 2.88, 0.55, 0.27, 13, accent, True, align=PP_ALIGN.RIGHT)
text(s, "验证链条：用户操作 → SDK / 服务端 → CallKit → RTC 远端输出", 0.78, 5.82,
     11.84, 0.45, 17, INK, True, align=PP_ALIGN.CENTER)

# 03 — architecture
s = base("一套用例，串起真实链路", "02 / 运行方式", 3, dark=True, time="约 1.5 分钟")
text(s, "Playwright 驱动独立浏览器上下文；每端使用不同 deviceId。", 0.7, 1.75,
     11.8, 0.46, 18, "B9C5D8")
for x, name, role in [(0.75, "A", "主叫"), (3.35, "B", "被叫"), (5.95, "C", "第三人 / 同账号另一设备")]:
    rect(s, x, 2.56, 2.25, 1.23, NAVY_2, True)
    circle(s, x+0.22, 2.82, 0.67, BLUE if name == "A" else TEAL if name == "B" else AMBER)
    text(s, name, x+0.22, 2.95, 0.67, 0.3, 18, WHITE, True, align=PP_ALIGN.CENTER)
    text(s, role, x+1.0, 2.93, 1.08, 0.35, 15, WHITE, True)
rect(s, 9.0, 2.56, 3.48, 1.23, "263856", True)
text(s, "真实 IM / 信令 / RTC", 9.28, 2.93, 2.9, 0.35, 18, WHITE, True,
     align=PP_ALIGN.CENTER)
for x1, x2 in [(3.0, 3.35), (5.6, 5.95), (8.2, 9.0)]:
    line(s, x1, 3.18, x2, 3.18, TEAL, 3)
rect(s, 0.75, 4.38, 11.73, 1.38, NAVY_2, True)
text(s, "断言与证据", 1.02, 4.67, 2.2, 0.35, 18, TEAL, True)
text(s, "服务端快照   ·   UI 状态   ·   浏览器事件 / trace   ·   远端音视频附件",
     3.15, 4.58, 8.95, 0.55, 18, WHITE)
text(s, "关键点：测试输入可控，服务端与 RTC 仍走真实链路。", 0.83, 6.22,
     11.8, 0.4, 16, "B9C5D8")

# 04 — video
s = base("真实视频：看见正确的人，也看见新画面", "03 / 视频验证", 4, time="约 2 分钟")
steps = [("01", "发送端", "Canvas 生成动态色块\n替换已发布视频轨"),
         ("02", "真实传输", "Agora RTC 发布、订阅\n远端解码 video"),
         ("03", "接收端", "采样远端像素\n检查解码帧持续增加")]
for i, (n, title, body) in enumerate(steps):
    x = 0.7 + i*4.18
    card(s, x, 1.9, 3.75, 1.88, title, body, [BLUE, TEAL, AMBER][i])
    text(s, n, x+3.0, 2.08, 0.42, 0.27, 12, [BLUE, TEAL, AMBER][i], True)
    if i < 2: line(s, x+3.77, 2.84, x+4.14, 2.84, TEAL, 3)
rect(s, 0.72, 4.25, 11.85, 1.55, WHITE, True)
text(s, "三人视频矩阵", 0.98, 4.47, 2.4, 0.34, 18, INK, True)
for i, (letter, color, name) in enumerate([("A", RED, "红"), ("B", TEAL, "绿"), ("C", BLUE, "蓝")]):
    x = 3.45 + i*2.88
    circle(s, x, 4.59, 0.56, color)
    text(s, letter, x+0.02, 4.72, 0.52, 0.23, 12, WHITE, True, align=PP_ALIGN.CENTER)
    text(s, f"{letter} 发送{name}色", x+0.7, 4.68, 1.75, 0.29, 14, INK, True)
text(s, "逐一验证 6 条接收链路，再轮换颜色；可发现黑屏、串流、冻结帧。",
     0.96, 5.14, 11.3, 0.38, 15, MUTED)
text(s, "证据：远端帧 PNG、像素比例 JSON、发送/接收轨道 ID", 0.83, 6.28,
     11.7, 0.31, 14, MUTED)

# 05 — audio
s = base("真实语音：让机器替我们“听”一次", "04 / 语音验证", 5, dark=True,
         time="约 2 分钟")
text(s, "固定中文句子：你好，现在开始语音通话测试。", 0.73, 1.7,
     11.9, 0.49, 21, WHITE, True)
names = [("发送", "音频送入主叫 RTC"), ("传输", "真实通话远端音轨"),
         ("录制", "被叫保存远端 WAV"), ("识别", "本地 ASR 得到文本")]
for i, (head, sub) in enumerate(names):
    x = 0.73 + i*3.15
    rect(s, x, 2.65, 2.82, 1.7, NAVY_2, True)
    circle(s, x+0.22, 2.97, 0.46, [BLUE, TEAL, AMBER, RED][i])
    text(s, str(i+1), x+0.23, 3.06, 0.44, 0.19, 11, WHITE, True, align=PP_ALIGN.CENTER)
    text(s, head, x+0.8, 2.92, 1.7, 0.36, 19, WHITE, True)
    text(s, sub, x+0.24, 3.55, 2.36, 0.55, 14, "B9C5D8")
    if i < 3: line(s, x+2.84, 3.48, x+3.13, 3.48, TEAL, 3)
rect(s, 0.74, 4.83, 11.83, 1.07, "263856", True)
text(s, "断言：录音长度 + 识别文本与原句按顺序匹配（容忍简繁差异）", 1.03,
     5.15, 11.15, 0.36, 18, WHITE, True)
text(s, "失败时保留接收端 WAV 与识别结果；比“麦克风图标亮着”更有说服力。",
     0.83, 6.31, 11.75, 0.35, 15, "B9C5D8")

# 06 — edge cases
s = base("更难的场景：竞态、掉线、多设备", "05 / 复杂场景", 6, time="约 1.5 分钟")
card(s, 0.72, 1.93, 3.82, 3.12, "同账号多端", "不同 deviceId 同时在线；接听后核对服务端 deviceUuid 与本机归属。", TEAL)
card(s, 4.76, 1.93, 3.82, 3.12, "响应超时", "分别模拟请求未到达与服务端已处理；用快照判定最终状态。", AMBER)
card(s, 8.80, 1.93, 3.82, 3.12, "成员故障", "异常掉线、视频发送中断与恢复；检查其他成员通话及远端媒体。", RED)
rect(s, 0.74, 5.53, 11.86, 0.8, INK, True)
text(s, "一个核心原则：以服务端状态和接收端结果共同判断，而不只看按钮。",
     1.02, 5.75, 11.3, 0.34, 18, WHITE, True)

# 07 — manual comparison
s = base("和手工测试相比，价值在哪里？", "06 / 测试价值", 7, dark=True,
         time="约 2 分钟")
headers = [("维度", 0.75, 2.45), ("手工测试", 3.05, 4.17), ("E2E 自动化", 7.3, 5.2)]
for label, x, w in headers:
    text(s, label, x, 1.77, w, 0.37, 17, TEAL, True)
rows = [
    ("协同", "人工协调多个账号、设备和操作时机", "独立浏览器上下文重复竞态时序"),
    ("判断", "靠观察画面、听声音", "验证远端像素、视频新帧、WAV 识别"),
    ("定位", "复现后再补日志", "快照、trace、媒体附件随失败保存"),
    ("回归", "每轮重复执行，容易漏边界", "同一批场景可持续重复运行"),
]
for i, (dim, manual, auto) in enumerate(rows):
    y = 2.32 + i*0.83
    rect(s, 0.73, y, 11.88, 0.71, NAVY_2 if i%2==0 else "1C2B47", True)
    text(s, dim, 0.9, y+0.17, 1.98, 0.33, 16, WHITE, True)
    text(s, manual, 3.06, y+0.13, 4.1, 0.43, 14, "B9C5D8")
    text(s, auto, 7.32, y+0.13, 5.04, 0.43, 14, WHITE)
text(s, "手工仍适合主观体验与真实硬件；自动化承担高频、可重复的客观回归。",
     0.84, 6.16, 11.74, 0.42, 16, "B9C5D8")

# 08 — debugging story
s = base("一次失败，如何缩短定位路径？", "07 / 排障实例", 8, time="约 1 分钟")
text(s, "以 WN-03 接听响应超时用例为例", 0.7, 1.68, 11.8, 0.42, 18, MUTED)
events = [("现象", "界面没有按预期收敛", BLUE),
          ("快照", "检查服务端参与者状态", TEAL),
          ("trace", "发现更早的 busy 1704", AMBER),
          ("结论", "先处理前置状态，再判断超时分支", RED)]
for i, (head, body, accent) in enumerate(events):
    x = 0.72 + i*3.14
    circle(s, x+0.07, 2.55, 0.57, accent)
    text(s, str(i+1), x+0.08, 2.68, 0.55, 0.25, 13, WHITE, True,
         align=PP_ALIGN.CENTER)
    if i < 3: line(s, x+0.67, 2.84, x+3.02, 2.84, LINE, 2)
    text(s, head, x+0.05, 3.39, 2.69, 0.39, 19, INK, True)
    text(s, body, x+0.05, 3.99, 2.73, 1.0, 15, MUTED, valign=MSO_ANCHOR.TOP)
rect(s, 0.74, 5.36, 11.85, 0.86, WHITE, True)
text(s, "收获：失败不只有“红灯”，还留下能区分环境、SDK、CallKit 与用例的问题证据。",
     1.0, 5.58, 11.3, 0.39, 17, INK, True)

# 09 — close
s = base("把“能连上”升级为“结果可信”", "08 / 总结", 9, dark=True, time="约 1 分钟")
for i, (head, sub, accent) in enumerate([
    ("状态正确", "服务端快照与 UI 一致", BLUE),
    ("归属正确", "接听设备与 RTC 参与者一致", TEAL),
    ("内容正确", "远端真的听见、看见新内容", AMBER),
]):
    x = 0.74 + i*4.18
    rect(s, x, 2.18, 3.78, 2.24, NAVY_2, True)
    circle(s, x+0.3, 2.56, 0.61, accent)
    text(s, f"0{i+1}", x+0.31, 2.7, 0.59, 0.25, 11, WHITE, True,
         align=PP_ALIGN.CENTER)
    text(s, head, x+1.08, 2.56, 2.34, 0.48, 23, WHITE, True)
    text(s, sub, x+0.3, 3.43, 3.12, 0.57, 16, "B9C5D8")
text(s, "下一步：把失败归因与环境就绪检查继续沉淀到回归流程。", 0.83,
     5.21, 11.65, 0.49, 19, WHITE, True)
text(s, "谢谢", 0.84, 6.17, 11.72, 0.43, 17, TEAL)

prs.save(OUT)
print(OUT)
